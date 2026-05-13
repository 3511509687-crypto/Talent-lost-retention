from pathlib import Path
import sys

sys.path.insert(0, str(Path("F:/app_bundle/pip_tmp")))

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("F:/app_bundle")
TEMPLATE = Path(r"D:\wx xiazai\WeChat Files\wxid_3cj0fkz1sbr822\FileStorage\File\2026-05\INF ALC Seminar S2W10 PPT.pptx")
OUT_DIR = ROOT / "presentation_prep"
ASSET_DIR = OUT_DIR / "04_ablation_oof_stacking" / "assets"
CROP_DIR = OUT_DIR / "s5_alc_template_cropped_assets"
PPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_4SLIDES_ALC_TEMPLATE_ACADEMIC_LAYOUT.pptx"
SCRIPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_4slides_ALC_template_academic_script.txt"

INK = (18, 31, 62)
MUTED = (72, 80, 98)
PURPLE = (113, 47, 211)
BLUE = (20, 75, 160)
TEAL = (20, 129, 106)
RED = (186, 75, 58)
LIGHT = (248, 249, 252)
LINE = (217, 223, 234)


def crop_near_white(src: Path, dst: Path, pad=14, threshold=250):
    im = Image.open(src).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg).convert("L")
    mask = diff.point(lambda p: 255 if p > 255 - threshold else 0)
    bbox = mask.getbbox()
    if bbox:
        left = max(bbox[0] - pad, 0)
        top = max(bbox[1] - pad, 0)
        right = min(bbox[2] + pad, im.size[0])
        bottom = min(bbox[3] + pad, im.size[1])
        im = im.crop((left, top, right, bottom))
    dst.parent.mkdir(parents=True, exist_ok=True)
    im.save(dst)
    return dst


def prepare_assets():
    names = [
        "ablation_metrics_comparison.png",
        "oof_stacking_flow.png",
        "topk_precision_recall_lift.png",
        "actual_vs_predicted_final_020_bins.png",
    ]
    return {name: crop_near_white(ASSET_DIR / name, CROP_DIR / name) for name in names}


def add_text(slide, x, y, w, h, text, size=12, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=11.5, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(4)
    return box


def add_card(slide, x, y, w, h, title, accent=PURPLE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*LIGHT)
    shape.line.color.rgb = RGBColor(*LINE)
    shape.line.width = Pt(0.8)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()
    add_text(slide, x + 0.15, y + 0.10, w - 0.25, 0.25, title, size=12.6, bold=True, color=INK)
    return shape


def add_metric(slide, x, y, w, label, value, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.55))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(*color)
    shape.line.width = Pt(1.0)
    add_text(slide, x + 0.05, y + 0.07, w - 0.10, 0.15, label, size=6.8, bold=True,
             color=(86, 92, 105), align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.05, y + 0.25, w - 0.10, 0.22, value, size=13.3, bold=True,
             color=color, align=PP_ALIGN.CENTER)


def add_image_fit(slide, path, x, y, w, h):
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    final_w = iw * scale
    final_h = ih * scale
    slide.shapes.add_picture(
        str(path),
        Inches(x + (w - final_w) / 2),
        Inches(y + (h - final_h) / 2),
        Inches(final_w),
        Inches(final_h),
    )


def add_title(slide, title, subtitle=""):
    add_text(slide, 0.33, 0.36, 8.95, 0.40, title, size=19.5, bold=True, color=INK)
    if subtitle:
        add_text(slide, 0.35, 0.82, 8.85, 0.25, subtitle, size=8.9, color=MUTED)


def clear_slide_content(slide):
    # Template slides are empty, but keep this defensive in case WeChat added hidden content.
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)


def build_deck():
    assets = prepare_assets()
    prs = Presentation(str(TEMPLATE))
    while len(prs.slides) < 4:
        prs.slides.add_slide(prs.slide_layouts[1])

    for i in range(4):
        clear_slide_content(prs.slides[i])

    # Slide 1
    s = prs.slides[0]
    add_title(s, "S5. Model Validation via Ablation and OOF Stacking",
              "This section establishes component necessity, leakage-controlled fusion, and operational ranking validity.")
    add_text(s, 0.55, 1.32, 4.20, 0.58,
             "The final ensemble is supported by three complementary evidence layers.",
             size=18.8, bold=True, color=INK)
    add_bullets(s, 0.60, 2.05, 4.10, 1.10, [
        "Component necessity: ablation quantifies marginal contribution.",
        "Fusion validity: OOF predictions reduce in-sample optimism.",
        "Operational value: top-k ranking prioritizes HR intervention resources.",
    ], size=10.4)
    for i, (label, value, color) in enumerate([
        ("OOF AUC", "0.9793", PURPLE),
        ("Test AUC", "0.9847", BLUE),
        ("Gap", "0.0054", TEAL),
        ("F1", "0.9248", RED),
    ]):
        add_metric(s, 0.60 + i * 1.05, 3.62, 0.92, label, value, color)
    add_image_fit(s, assets["oof_stacking_flow.png"], 5.02, 1.36, 4.50, 1.45)
    add_card(s, 5.12, 3.33, 4.02, 0.98, "Analytical progression", BLUE)
    add_text(s, 5.35, 3.78, 3.52, 0.25,
             "Component evidence -> fair fusion -> deployable risk ranking",
             size=11.7, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

    # Slide 2
    s = prs.slides[1]
    add_title(s, "Ablation Study: Quantifying Marginal Module Contribution",
              "Each major component is removed or replaced to estimate its incremental effect on predictive performance.")
    add_image_fit(s, assets["ablation_metrics_comparison.png"], 0.38, 1.20, 6.70, 3.82)
    add_card(s, 7.20, 1.30, 2.28, 0.96, "Estimation principle", PURPLE)
    add_text(s, 7.34, 1.77, 2.00, 0.22, "Contribution_j = AUC_full - AUC_without_j",
             size=8.5, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    add_card(s, 7.20, 2.55, 2.28, 1.78, "Interpretation", BLUE)
    add_bullets(s, 7.37, 3.01, 1.96, 0.98, [
        "TF-IDF provides a lexical baseline.",
        "Sentence-BERT contributes semantic policy signals.",
        "LR, ET and LGB isolate linear, bagging-style, and boosting-style effects.",
    ], size=8.1)
    add_text(s, 0.45, 5.04, 8.95, 0.18,
             "Conclusion: the reported performance is supported by controlled component-level evidence.",
             size=8.5, bold=True, color=(42, 49, 63))

    # Slide 3
    s = prs.slides[2]
    add_title(s, "OOF Stacking: Leakage-Controlled Ensemble Learning",
              "The meta learner is trained on out-of-fold base-model predictions rather than in-sample fitted outputs.")
    add_image_fit(s, assets["oof_stacking_flow.png"], 0.55, 1.20, 8.90, 1.55)
    cards = [
        (0.65, "Out-of-fold score", "p_i^OOF = f_-k(x_i)", PURPLE),
        (3.72, "Meta-feature vector", "z_i = [p_LR, p_ET, p_LGB, mean, std, disagreement]", BLUE),
        (6.79, "Stacked prediction", "P(y_i=1|x_i) = sigma(w^T z_i + b)", TEAL),
    ]
    for x, title, formula, color in cards:
        add_card(s, x, 3.05, 2.55, 0.95, title, color)
        add_text(s, x + 0.14, 3.48, 2.28, 0.24, formula, size=8.9, bold=True,
                 color=color, align=PP_ALIGN.CENTER)
    add_text(s, 0.78, 4.43, 8.3, 0.43,
             "Compared with uniform averaging, the meta learner estimates context-specific model reliability and treats model disagreement as uncertainty information.",
             size=9.9, bold=True, color=(42, 49, 66), align=PP_ALIGN.CENTER)

    # Slide 4
    s = prs.slides[3]
    add_title(s, "Operational Evaluation: Risk Ranking for HR Intervention",
              "The validated model is translated into a prioritized Top-20% employee risk list.")
    add_metric(s, 0.58, 1.10, 1.05, "Top 20% Precision", "0.9700", TEAL)
    add_metric(s, 1.78, 1.10, 1.05, "Top 20% Recall", "0.8151", BLUE)
    add_metric(s, 2.98, 1.10, 1.05, "Top 20% Lift", "4.0756", PURPLE)
    add_image_fit(s, assets["topk_precision_recall_lift.png"], 0.48, 1.88, 4.75, 3.05)
    add_text(s, 5.55, 1.10, 3.8, 0.25, "Calibration evidence: 20-bin reliability view",
             size=11.8, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_image_fit(s, assets["actual_vs_predicted_final_020_bins.png"], 5.37, 1.55, 4.05, 2.52)
    add_card(s, 5.72, 4.35, 3.34, 0.78, "Transition to S6", RED)
    add_text(s, 5.92, 4.72, 2.94, 0.20,
             "SHAP provides individual-level attribution for each risk flag.",
             size=10.2, bold=True, color=RED, align=PP_ALIGN.CENTER)

    prs.save(PPT_OUT)


def write_script():
    script = """Slide 1.
S4 established LightGBM as a strong nonlinear predictor. This section presents the validation logic behind the final ensemble. The design is supported by ablation analysis, out-of-fold stacking, and HR-oriented ranking evaluation.

Slide 2.
Ablation analysis removes or replaces one component and then measures the resulting performance change. We compare TF-IDF with Sentence-BERT, evaluate LR, ET and LightGBM separately, and test full-model variants without key components. This provides component-level evidence rather than only reporting the best final score.

Slide 3.
OOF means out-of-fold. Each training sample receives a prediction from a model that did not train on that sample. These leakage-controlled predictions become inputs for the logistic meta learner. Compared with simple averaging, stacking can assign different reliability to LR, ET and LightGBM, while also using disagreement as uncertainty information.

Slide 4.
The full model reaches OOF AUC 0.9793 and Test AUC 0.9847, with only a 0.0054 gap. For HR action, the Top 20 percent risk group reaches Precision 0.9700, Recall 0.8151, and Lift 4.0756. The model therefore produces a focused intervention list. The next section uses SHAP to explain individual risk drivers."""
    SCRIPT_OUT.write_text(script, encoding="utf-8")


if __name__ == "__main__":
    build_deck()
    write_script()
    print(f"Wrote {PPT_OUT}")
    print(f"Wrote {SCRIPT_OUT}")
