from pathlib import Path
import sys

sys.path.insert(0, str(Path("F:/app_bundle/pip_tmp")))

from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("F:/app_bundle")
OUT_DIR = ROOT / "presentation_prep"
ASSET_DIR = OUT_DIR / "04_ablation_oof_stacking" / "assets"
PPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_4SLIDES_POLISHED_XJTLU_style.pptx"
SCRIPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_4slides_script.txt"
FOOTER_IMG = OUT_DIR / "s5_xjtlu_footer_gradient.png"
CROP_DIR = OUT_DIR / "s5_four_slide_cropped_assets"

PURPLE = (96, 52, 210)
BLUE = (21, 73, 170)
TEAL = (18, 117, 92)
RED = (185, 71, 54)
INK = (15, 28, 61)
MUTED = (78, 85, 101)
LIGHT = (248, 249, 252)
LINE = (218, 222, 232)


def make_footer(path: Path, width=1920, height=72):
    left = (200, 0, 255)
    right = (0, 34, 92)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    for x in range(width):
        t = x / max(width - 1, 1)
        color = tuple(int(left[i] * (1 - t) + right[i] * t) for i in range(3))
        draw.line([(x, 0), (x, height)], fill=color)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def crop_near_white(src: Path, dst: Path, pad=18, threshold=250):
    im = Image.open(src).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg).convert("L")
    mask = diff.point(lambda p: 255 if p > 255 - threshold else 0)
    bbox = mask.getbbox()
    if bbox:
        left = max(bbox[0] - pad, 0)
        top = max(bbox[1] - pad, 0)
        right = min(bbox[2] + pad, im.size[0])
        bottom = min(bbox[3] + pad, im.size[1])
        im = im.crop((left, top, right, bottom))
    dst.parent.mkdir(parents=True, exist_ok=True)
    im.save(dst)
    return dst


def prepare_assets():
    names = [
        "ablation_metrics_comparison.png",
        "oof_stacking_flow.png",
        "topk_precision_recall_lift.png",
        "actual_vs_predicted_final_020_bins.png",
    ]
    return {name: crop_near_white(ASSET_DIR / name, CROP_DIR / name) for name in names}


def add_text(slide, x, y, w, h, text, size=12, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=13, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(5)
    return box


def add_card(slide, x, y, w, h, title, accent=PURPLE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*LIGHT)
    shape.line.color.rgb = RGBColor(*LINE)
    shape.line.width = Pt(0.9)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()
    add_text(slide, x + 0.18, y + 0.13, w - 0.34, 0.35, title, size=15, bold=True, color=INK)
    return shape


def add_metric(slide, x, y, w, label, value, color):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.62))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.line.color.rgb = RGBColor(*color)
    box.line.width = Pt(1.1)
    add_text(slide, x + 0.08, y + 0.08, w - 0.16, 0.17, label, size=7.5, bold=True, color=(88, 94, 108),
             align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.08, y + 0.27, w - 0.16, 0.27, value, size=15, bold=True, color=color,
             align=PP_ALIGN.CENTER)


def add_image_fit(slide, path, x, y, w, h):
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    final_w = iw * scale
    final_h = ih * scale
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - final_w) / 2),
        Inches(y + (h - final_h) / 2),
        Inches(final_w),
        Inches(final_h),
    )


def add_base(slide, section, title, subtitle=""):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    slide.shapes.add_picture(str(FOOTER_IMG), Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
    side = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.72), Inches(0.14), Inches(0.64))
    side.fill.solid()
    side.fill.fore_color.rgb = RGBColor(203, 0, 255)
    side.line.fill.background()
    add_text(slide, 0.45, 0.30, 1.0, 0.24, section, size=9, bold=True, color=PURPLE)
    add_text(slide, 0.45, 0.52, 8.9, 0.44, title, size=24, bold=True, color=INK)
    if subtitle:
        add_text(slide, 0.46, 0.96, 10.7, 0.27, subtitle, size=10.8, color=MUTED)
    add_text(slide, 10.55, 7.18, 2.25, 0.22, "XJTLU | ACADEMIC LITERACIES CENTRE",
             size=8.5, bold=True, color=(255, 255, 255), align=PP_ALIGN.RIGHT)


def build_deck():
    make_footer(FOOTER_IMG)
    assets = prepare_assets()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "S5 / Validation", "Ablation + OOF Stacking",
             "Goal: prove the final model is designed, validated, and useful for HR action.")
    add_text(s, 0.65, 1.65, 6.0, 0.78,
             "How do we know the final model is not just a random stack of algorithms?",
             size=26, bold=True, color=INK)
    add_bullets(s, 0.70, 2.62, 5.4, 1.45, [
        "Ablation tests component contribution.",
        "OOF stacking gives leakage-controlled training signals.",
        "Top-20% ranking converts prediction into an action list.",
    ], size=14)
    for i, (label, value, color) in enumerate([
        ("OOF AUC", "0.9793", PURPLE),
        ("Test AUC", "0.9847", BLUE),
        ("OOF/Test Gap", "0.0054", TEAL),
        ("Test F1", "0.9248", RED),
    ]):
        add_metric(s, 0.72 + i * 1.36, 4.55, 1.18, label, value, color)
    add_image_fit(s, assets["oof_stacking_flow.png"], 6.55, 1.55, 6.15, 2.0)
    add_card(s, 6.72, 3.86, 5.55, 1.36, "Section logic", BLUE)
    add_text(s, 6.94, 4.35, 5.06, 0.45,
             "Component proof -> fair fusion -> ranked HR intervention list",
             size=18, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

    # Slide 2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "S5 / Ablation", "Ablation Study: Does Each Module Matter?",
             "Remove or replace one component and measure the performance change.")
    add_image_fit(s, assets["ablation_metrics_comparison.png"], 0.55, 1.48, 8.05, 4.85)
    add_card(s, 8.9, 1.50, 3.85, 1.25, "Evidence logic", PURPLE)
    add_text(s, 9.15, 2.05, 3.35, 0.34, "Delta AUC_j = AUC_full - AUC_without_j",
             size=13.5, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    add_card(s, 8.9, 3.03, 3.85, 2.25, "What it proves", BLUE)
    add_bullets(s, 9.12, 3.55, 3.30, 1.35, [
        "TF-IDF is the lexical baseline.",
        "Sentence-BERT adds semantic policy matching.",
        "LR, ET and LGB are tested alone and inside the full model.",
    ], size=10.7)
    add_text(s, 0.58, 6.70, 11.4, 0.20,
             "Takeaway: the final score is supported by controlled comparisons, not only by one best-performing run.",
             size=9.3, bold=True, color=(41, 48, 63))

    # Slide 3
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "S5 / OOF Stacking", "OOF Stacking: Why Not Simple Average?",
             "Each training row is scored by models that did not see that row, then a meta learner combines the probabilities.")
    add_image_fit(s, assets["oof_stacking_flow.png"], 0.58, 1.42, 12.0, 2.25)
    add_card(s, 0.72, 4.05, 3.72, 1.35, "OOF prediction", PURPLE)
    add_text(s, 0.95, 4.62, 3.25, 0.25, "p_i^OOF = f_-k(x_i)", size=15, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    add_card(s, 4.80, 4.05, 3.72, 1.35, "Meta features", BLUE)
    add_text(s, 5.02, 4.52, 3.25, 0.46, "z_i = [p_LR, p_ET, p_LGB, mean, std, disagreement]",
             size=11.6, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_card(s, 8.88, 4.05, 3.72, 1.35, "Final probability", TEAL)
    add_text(s, 9.10, 4.62, 3.25, 0.25, "P_final = sigmoid(w^T z_i + b)",
             size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_text(s, 1.02, 5.92, 11.0, 0.36,
             "Simple average gives equal trust. Stacking learns when LR, ET or LightGBM should matter more, and uses disagreement as uncertainty information.",
             size=12.2, bold=True, color=(42, 49, 66), align=PP_ALIGN.CENTER)

    # Slide 4
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(s, "S5 / HR Action", "From Validated Model to HR Action List",
             "The final output is not only a probability, but a ranked intervention list.")
    add_metric(s, 0.70, 1.35, 1.35, "Top 20% Precision", "0.9700", TEAL)
    add_metric(s, 2.25, 1.35, 1.35, "Top 20% Recall", "0.8151", BLUE)
    add_metric(s, 3.80, 1.35, 1.35, "Top 20% Lift", "4.0756", PURPLE)
    add_image_fit(s, assets["topk_precision_recall_lift.png"], 0.58, 2.18, 6.2, 4.18)
    add_text(s, 7.15, 1.40, 5.15, 0.28, "Final calibration view: 20 bins",
             size=13.4, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_image_fit(s, assets["actual_vs_predicted_final_020_bins.png"], 7.05, 1.82, 5.45, 3.10)
    add_card(s, 7.26, 5.25, 4.92, 1.06, "Transition to S6", RED)
    add_text(s, 7.55, 5.75, 4.34, 0.28,
             "After validation and ranking, SHAP answers why an employee is flagged.",
             size=13.2, bold=True, color=RED, align=PP_ALIGN.CENTER)

    prs.save(PPT_OUT)


def write_script():
    script = """Slide 1.
S4 showed that LightGBM is strong. My section asks a validation question: how do we know the final model is not just a random stack of algorithms? We answer this with ablation, OOF stacking, and HR ranking.

Slide 2.
Ablation means removing or replacing one component and checking the performance change. We compare TF-IDF with Sentence-BERT, test LR, ET and LightGBM separately, and test the full model without key components. This proves contribution, not only final performance.

Slide 3.
OOF means out-of-fold. Each training sample receives a prediction from a model that did not train on that sample. These fair predictions become inputs for the meta learner. This is better than simple average because LR, ET and LightGBM are reliable in different situations, and disagreement itself is useful.

Slide 4.
The full model reaches OOF AUC 0.9793 and Test AUC 0.9847, with only a 0.0054 gap. For HR action, the Top 20 percent list reaches Precision 0.9700, Recall 0.8151, and Lift 4.0756. So the model produces a focused action list. Next, SHAP explains why each employee is flagged."""
    SCRIPT_OUT.write_text(script, encoding="utf-8")


if __name__ == "__main__":
    build_deck()
    write_script()
    print(f"Wrote {PPT_OUT}")
    print(f"Wrote {SCRIPT_OUT}")
