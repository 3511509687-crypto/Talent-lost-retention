from pathlib import Path
import sys

sys.path.insert(0, str(Path("F:/app_bundle/pip_tmp")))

from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE


ROOT = Path("F:/app_bundle")
OUT_DIR = ROOT / "presentation_prep"
ASSET_DIR = OUT_DIR / "04_ablation_oof_stacking" / "assets"

PPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_ONE_PAGE_POLISHED_XJTLU_style.pptx"
SCRIPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_one_page_script.txt"
FOOTER_IMG = OUT_DIR / "s5_xjtlu_footer_gradient.png"
CROP_DIR = OUT_DIR / "s5_one_page_cropped_assets"


def make_footer(path: Path, width=1920, height=72):
    left = (200, 0, 255)
    right = (0, 34, 92)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    for x in range(width):
        t = x / max(width - 1, 1)
        color = tuple(int(left[i] * (1 - t) + right[i] * t) for i in range(3))
        draw.line([(x, 0), (x, height)], fill=color)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def crop_near_white(src: Path, dst: Path, pad=16, threshold=245):
    im = Image.open(src).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg).convert("L")
    mask = diff.point(lambda p: 255 if p > 255 - threshold else 0)
    bbox = mask.getbbox()
    if not bbox:
        im.save(dst)
        return dst
    left = max(bbox[0] - pad, 0)
    top = max(bbox[1] - pad, 0)
    right = min(bbox[2] + pad, im.size[0])
    bottom = min(bbox[3] + pad, im.size[1])
    im.crop((left, top, right, bottom)).save(dst)
    return dst


def prepare_cropped_assets():
    CROP_DIR.mkdir(parents=True, exist_ok=True)
    names = [
        "ablation_metrics_comparison.png",
        "oof_stacking_flow.png",
        "topk_precision_recall_lift.png",
        "actual_vs_predicted_final_020_bins.png",
    ]
    return {
        name: crop_near_white(ASSET_DIR / name, CROP_DIR / name)
        for name in names
    }


def add_textbox(slide, x, y, w, h, text, size=12, bold=False, color=(20, 25, 35),
                align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=9.2):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, text in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(55, 61, 74)
        p.space_after = Pt(1.5)
    return box


def add_card(slide, x, y, w, h, title, accent):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(248, 249, 252)
    shape.line.color.rgb = RGBColor(218, 222, 232)
    shape.line.width = Pt(0.8)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()
    add_textbox(slide, x + 0.18, y + 0.12, w - 0.3, 0.28, title, size=13, bold=True, color=(17, 31, 60))
    return shape


def add_image_fit(slide, path, x, y, w, h):
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    final_w = iw * scale
    final_h = ih * scale
    left = x + (w - final_w) / 2
    top = y + (h - final_h) / 2
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(final_w), Inches(final_h))


def add_metric_chip(slide, x, y, w, label, value, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.52))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(*color)
    shape.line.width = Pt(1.2)
    add_textbox(slide, x + 0.08, y + 0.06, w - 0.16, 0.16, label, size=6.8, bold=True, color=(91, 96, 110),
                align=PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.08, y + 0.22, w - 0.16, 0.23, value, size=13, bold=True, color=color,
                align=PP_ALIGN.CENTER)


def build_ppt():
    make_footer(FOOTER_IMG)
    assets = prepare_cropped_assets()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    slide.shapes.add_picture(str(FOOTER_IMG), Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
    side = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.72), Inches(0.14), Inches(0.64))
    side.fill.solid()
    side.fill.fore_color.rgb = RGBColor(203, 0, 255)
    side.line.fill.background()

    add_textbox(slide, 0.45, 0.30, 5.7, 0.42, "S5. Ablation + OOF Stacking", size=23, bold=True, color=(15, 28, 61))
    add_textbox(slide, 0.46, 0.73, 7.8, 0.26,
                "Validation question: which components matter, and can the final ranking generalize?",
                size=10.2, color=(80, 86, 100))
    add_textbox(slide, 10.55, 7.18, 2.25, 0.22, "XJTLU | ACADEMIC LITERACIES CENTRE",
                size=8.5, bold=True, color=(255, 255, 255), align=PP_ALIGN.RIGHT)

    chips = [
        ("OOF AUC", "0.9793", (96, 52, 210)),
        ("Test AUC", "0.9847", (21, 73, 170)),
        ("Gap", "0.0054", (44, 128, 115)),
        ("Top 20% Lift", "4.0756", (18, 117, 92)),
    ]
    for i, (label, value, color) in enumerate(chips):
        add_metric_chip(slide, 7.95 + i * 1.18, 0.30, 1.06, label, value, color)

    left_x, left_y, left_w, left_h = 0.45, 1.10, 7.55, 5.84
    right_x, right_y, right_w, right_h = 8.23, 1.10, 4.67, 5.84

    add_card(slide, left_x, left_y, left_w, left_h, "Validation design: ablation + OOF stacking", (97, 65, 201))
    add_textbox(slide, left_x + 0.22, left_y + 0.45, left_w - 0.44, 0.16,
                "OOF gives each row a prediction from models that did not train on it; ablation then tests what each module contributes.",
                size=7.8, color=(72, 78, 92))
    add_image_fit(slide, assets["oof_stacking_flow.png"], left_x + 0.15, left_y + 0.72, left_w - 0.30, 1.36)
    formula = "p_i^OOF = f_-k(x_i)   |   P_final = sigmoid(w^T z_i + b)"
    add_textbox(slide, left_x + 0.24, left_y + 2.26, left_w - 0.48, 0.18, formula,
                size=8.6, bold=True, color=(20, 51, 105), align=PP_ALIGN.CENTER)
    add_textbox(slide, left_x + 0.22, left_y + 2.56, left_w - 0.44, 0.20,
                "Ablation evidence: compare full model with baselines and removed components",
                size=8.5, bold=True, color=(44, 50, 66), align=PP_ALIGN.CENTER)
    add_image_fit(slide, assets["ablation_metrics_comparison.png"], left_x + 0.15, left_y + 2.80, left_w - 0.30, 3.03)

    add_card(slide, right_x, right_y, right_w, right_h, "HR action: ranked high-risk list", (25, 142, 105))
    add_textbox(slide, right_x + 0.22, right_y + 0.46, right_w - 0.44, 0.22,
                "Top 20%: Precision 0.9700 | Recall 0.8151 | Lift 4.0756",
                size=9.0, bold=True, color=(18, 117, 92), align=PP_ALIGN.CENTER)
    add_image_fit(slide, assets["topk_precision_recall_lift.png"], right_x + 0.11, right_y + 0.78, right_w - 0.22, 2.66)
    add_textbox(slide, right_x + 0.22, right_y + 3.42, right_w - 0.44, 0.18,
                "Final calibration view: 20 bins",
                size=8.1, bold=True, color=(65, 72, 88), align=PP_ALIGN.CENTER)
    add_image_fit(slide, assets["actual_vs_predicted_final_020_bins.png"], right_x + 0.11, right_y + 3.66, right_w - 0.22, 2.16)

    add_textbox(slide, 0.48, 6.88, 11.7, 0.20,
                "Takeaway: component contribution is tested by ablation, generalization is protected by OOF stacking, and the output becomes a focused HR action list.",
                size=8.3, bold=True, color=(37, 43, 57))

    prs.save(PPT_OUT)


def write_script():
    script = """S4 showed that LightGBM is strong. This slide checks whether the final system is really justified.

We validate it in two ways. First, ablation study: we remove or replace one component, such as TF-IDF versus Sentence-BERT, or LR, ET and LightGBM, and then compare the change in performance. This tells us which modules contribute to the final result, instead of only reporting a high score.

Second, we use OOF stacking. OOF means out-of-fold: every training sample receives a prediction from a model that did not train on that sample. These fair predictions are then used by a logistic meta learner to combine LR, ET and LightGBM. This is better than simple average because different models are reliable in different cases, and disagreement itself is useful information.

The evidence is stable. Our full model reaches OOF AUC 0.9793 and Test AUC 0.9847, with only a 0.0054 gap. It also reaches Test F1 0.9248.

For HR action, the Top 20 percent risk list has Precision 0.9700, Recall 0.8151 and Lift 4.0756. So the model does not only predict attrition; it produces a focused action list. Next, SHAP explains why each employee is flagged."""
    SCRIPT_OUT.write_text(script, encoding="utf-8")


if __name__ == "__main__":
    build_ppt()
    write_script()
    print(f"Wrote {PPT_OUT}")
    print(f"Wrote {SCRIPT_OUT}")
