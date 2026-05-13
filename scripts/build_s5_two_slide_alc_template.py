from pathlib import Path
import sys

sys.path.insert(0, str(Path("F:/app_bundle/pip_tmp")))

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("F:/app_bundle")
TEMPLATE = Path(r"D:\wx xiazai\WeChat Files\wxid_3cj0fkz1sbr822\FileStorage\File\2026-05\INF ALC Seminar S2W10 PPT.pptx")
OUT_DIR = ROOT / "presentation_prep"
ASSET_DIR = OUT_DIR / "04_ablation_oof_stacking" / "assets"
CROP_DIR = OUT_DIR / "s5_two_slide_cropped_assets"
PPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_2SLIDES_ALC_TEMPLATE_ACADEMIC.pptx"
SCRIPT_OUT = OUT_DIR / "S5_Ablation_OOF_Stacking_2slides_ALC_template_academic_script.txt"

INK = (18, 31, 62)
MUTED = (72, 80, 98)
PURPLE = (113, 47, 211)
BLUE = (20, 75, 160)
TEAL = (20, 129, 106)
RED = (186, 75, 58)
LIGHT = (248, 249, 252)
LINE = (217, 223, 234)


def crop_near_white(src: Path, dst: Path, pad=14, threshold=250):
    im = Image.open(src).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg).convert("L")
    mask = diff.point(lambda p: 255 if p > 255 - threshold else 0)
    bbox = mask.getbbox()
    if bbox:
        left = max(bbox[0] - pad, 0)
        top = max(bbox[1] - pad, 0)
        right = min(bbox[2] + pad, im.size[0])
        bottom = min(bbox[3] + pad, im.size[1])
        im = im.crop((left, top, right, bottom))
    dst.parent.mkdir(parents=True, exist_ok=True)
    im.save(dst)
    return dst


def prepare_assets():
    names = [
        "ablation_metrics_comparison.png",
        "oof_stacking_flow.png",
        "topk_precision_recall_lift.png",
        "actual_vs_predicted_final_020_bins.png",
    ]
    return {name: crop_near_white(ASSET_DIR / name, CROP_DIR / name) for name in names}


def delete_slide(prs, index):
    sld_id_lst = prs.slides._sldIdLst
    sld_ids = list(sld_id_lst)
    r_id = sld_ids[index].rId
    prs.part.drop_rel(r_id)
    sld_id_lst.remove(sld_ids[index])


def clear_slide_content(slide):
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)


def add_text(slide, x, y, w, h, text, size=12, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=9.2, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(3.5)
    return box


def add_card(slide, x, y, w, h, title, accent=PURPLE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*LIGHT)
    shape.line.color.rgb = RGBColor(*LINE)
    shape.line.width = Pt(0.8)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()
    add_text(slide, x + 0.14, y + 0.10, w - 0.22, 0.23, title, size=11.2, bold=True, color=INK)
    return shape


def add_metric(slide, x, y, w, label, value, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.50))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(*color)
    shape.line.width = Pt(0.95)
    add_text(slide, x + 0.05, y + 0.06, w - 0.10, 0.14, label, size=6.2, bold=True,
             color=(86, 92, 105), align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.05, y + 0.22, w - 0.10, 0.22, value, size=12.3, bold=True,
             color=color, align=PP_ALIGN.CENTER)


def add_image_fit(slide, path, x, y, w, h):
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    final_w = iw * scale
    final_h = ih * scale
    slide.shapes.add_picture(
        str(path),
        Inches(x + (w - final_w) / 2),
        Inches(y + (h - final_h) / 2),
        Inches(final_w),
        Inches(final_h),
    )


def add_title(slide, title, subtitle=""):
    add_text(slide, 0.33, 0.36, 8.95, 0.40, title, size=19.2, bold=True, color=INK)
    if subtitle:
        add_text(slide, 0.35, 0.82, 8.85, 0.25, subtitle, size=8.8, color=MUTED)


def build_deck():
    assets = prepare_assets()
    prs = Presentation(str(TEMPLATE))
    while len(prs.slides) < 2:
        prs.slides.add_slide(prs.slide_layouts[1])
    while len(prs.slides) > 2:
        delete_slide(prs, len(prs.slides) - 1)

    for slide in prs.slides:
        clear_slide_content(slide)

    # Slide 1: framework + ablation evidence
    s = prs.slides[0]
    add_title(
        s,
        "S5. Validation Framework and Ablation Evidence",
        "The final ensemble is supported by component-level tests before model fusion is evaluated.",
    )
    add_image_fit(s, assets["ablation_metrics_comparison.png"], 0.38, 1.23, 6.35, 3.66)
    add_card(s, 6.95, 1.22, 2.55, 1.18, "Validation framework", PURPLE)
    add_bullets(s, 7.13, 1.65, 2.15, 0.58, [
        "Component necessity",
        "Fusion validity",
        "HR actionability",
    ], size=8.5)
    add_card(s, 6.95, 2.62, 2.55, 1.18, "Ablation principle", BLUE)
    add_text(s, 7.15, 3.10, 2.15, 0.24, "Contribution_j = AUC_full - AUC_without_j",
             size=8.0, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_card(s, 6.95, 4.02, 2.55, 0.70, "Key interpretation", TEAL)
    add_text(s, 7.12, 4.37, 2.20, 0.15,
             "Semantic features and heterogeneous base models are retained because their removal weakens the validation evidence.",
             size=7.5, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_text(s, 0.44, 5.04, 8.80, 0.18,
             "Conclusion: ablation converts final-model performance into controlled evidence of module contribution.",
             size=8.2, bold=True, color=(42, 49, 63))

    # Slide 2: OOF stacking + HR evaluation
    s = prs.slides[1]
    add_title(
        s,
        "OOF Stacking and HR-Oriented Ranking Evaluation",
        "Out-of-fold fusion controls training optimism, while Top-20% ranking translates predictions into intervention priorities.",
    )
    add_image_fit(s, assets["oof_stacking_flow.png"], 0.43, 1.14, 5.10, 1.16)
    add_card(s, 0.47, 2.48, 2.45, 0.82, "Stacking formulation", PURPLE)
    add_text(s, 0.62, 2.86, 2.13, 0.18, "P(y=1|x) = sigma(w^T z_i + b)",
             size=8.5, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    add_card(s, 3.08, 2.48, 2.45, 0.82, "Meta-feature vector", BLUE)
    add_text(s, 3.22, 2.82, 2.17, 0.24, "z_i = [p_LR, p_ET, p_LGB, mean, std, disagreement]",
             size=7.2, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

    for i, (label, value, color) in enumerate([
        ("OOF AUC", "0.9793", PURPLE),
        ("Test AUC", "0.9847", BLUE),
        ("Gap", "0.0054", TEAL),
        ("F1", "0.9248", RED),
    ]):
        add_metric(s, 5.78 + i * 0.92, 1.16, 0.78, label, value, color)
    add_image_fit(s, assets["topk_precision_recall_lift.png"], 5.75, 1.92, 3.75, 2.42)
    add_text(s, 5.92, 4.38, 3.35, 0.18,
             "Top 20%: Precision 0.9700 | Recall 0.8151 | Lift 4.0756",
             size=7.9, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_text(s, 0.55, 3.62, 4.70, 0.22,
             "Compared with uniform averaging, the meta learner estimates context-specific reliability across LR, ET and LightGBM.",
             size=8.4, bold=True, color=(42, 49, 66), align=PP_ALIGN.CENTER)
    add_image_fit(s, assets["actual_vs_predicted_final_020_bins.png"], 0.95, 4.00, 3.90, 0.98)
    add_text(s, 5.80, 4.88, 3.62, 0.16,
             "Transition: SHAP provides individual-level attribution for the generated risk flags.",
             size=7.6, bold=True, color=RED, align=PP_ALIGN.CENTER)

    prs.save(PPT_OUT)


def write_script():
    script = """Slide 1.
S4 established LightGBM as a strong nonlinear predictor. This section compresses the validation logic into two evidence layers. First, the ablation study evaluates component necessity by removing or replacing key modules and observing the performance change. The comparison covers lexical TF-IDF, semantic Sentence-BERT features, and the LR, ET and LightGBM base models. This turns the final score into component-level evidence rather than a single isolated result.

Slide 2.
The second layer evaluates fusion and operational use. OOF stacking trains the meta learner on predictions produced by models that did not see the corresponding training samples, reducing in-sample optimism. The meta learner combines LR, ET and LightGBM probabilities, along with average, spread and disagreement features. The full model reaches OOF AUC 0.9793 and Test AUC 0.9847, with a 0.0054 gap. For HR action, the Top 20 percent risk group reaches Precision 0.9700, Recall 0.8151 and Lift 4.0756. The next section uses SHAP to explain individual risk drivers."""
    SCRIPT_OUT.write_text(script, encoding="utf-8")


if __name__ == "__main__":
    build_deck()
    write_script()
    print(f"Wrote {PPT_OUT}")
    print(f"Wrote {SCRIPT_OUT}")
